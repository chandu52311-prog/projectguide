import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def set_cell_background(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))


def format_inductance(val: float) -> str:
    if val <= 0:
        return "0 H"
    if val < 1e-3:
        return f"{val * 1e6:.2f} µH"
    if val < 1.0:
        return f"{val * 1e3:.3f} mH"
    return f"{val:.4f} H"


def format_capacitance(val: float) -> str:
    if val <= 0:
        return "0 F"
    if val < 1e-6:
        return f"{val * 1e9:.2f} nF"
    if val < 1e-3:
        return f"{val * 1e6:.2f} µF"
    return f"{val * 1e3:.2f} mF"


def _table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, h in zip(t.rows[0].cells, headers):
        cell.text = h
        run = cell.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "0F172A")
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = str(v)
            c.paragraphs[0].runs[0].font.size = Pt(9.5)
    doc.add_paragraph()


def generate_project_report(branch: str, level: str, ratings: dict, lcl: dict,
                            custom_notes: str = "", provenance: dict = None,
                            novelty_type: str = "LLCL_EXTENSION") -> io.BytesIO:
    provenance = provenance or {}
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(36)
    r = title.add_run(f"SIMULATION SPECIFICATION REPORT\n{level.upper()} PROJECT SUPPORT DOCUMENT")
    r.font.size, r.font.bold, r.font.name = Pt(18), True, "Calibri"
    r.font.color.rgb = RGBColor(15, 23, 42)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f"Discipline: {branch} Engineering\nExtension: {novelty_type}\n"
                     "Design basis: LCL/LLCL grid filter sizing (IEEE Std 519 harmonic context)")
    sr.font.size, sr.font.italic = Pt(11), True
    doc.add_page_break()

    doc.add_heading("Chapter 1: System Ratings", level=1)
    doc.add_paragraph("Values marked 'default' were NOT found in the base paper and must be replaced "
                      "with the paper's actual ratings before the results are used.")
    prov = lambda k: provenance.get(k, "n/a")
    _table(doc, ["Parameter", "Value", "Source"], [
        ("DC bus voltage (Vdc)", f"{ratings['v_dc']:.1f} V", prov("v_dc")),
        ("Grid phase voltage (Vrms)", f"{ratings['v_grid']:.1f} V", prov("v_grid")),
        ("Switching frequency (fsw)", f"{ratings['f_sw'] / 1000:.2f} kHz", prov("f_sw")),
        ("Rated power (P)", f"{ratings['p_rated'] / 1000:.2f} kW", prov("p_rated")),
        ("Grid frequency", f"{ratings.get('f_grid', 50.0):.0f} Hz", prov("f_grid")),
    ])

    doc.add_heading("Chapter 2: Filter Sizing and Analytical Check", level=1)
    doc.add_paragraph(
        "Inverter-side inductance follows from the allowed peak-to-peak current ripple; the capacitor is "
        "limited to 5 % of base capacitance; the grid-side inductance follows from the required ripple "
        "attenuation ig/ii = 1/|1 + r(1 - L1·Cf·ωsw²)|, r = L2/L1; damping Rd = 1/(3·ωres·Cf).")
    v = lcl["verification"]
    _table(doc, ["Component", "Computed", "Check"], [
        ("L1 (inverter side)", format_inductance(lcl["L1_inverter_H"]), "ripple ≤ 15 % of Ipk"),
        ("L2 (grid side)", format_inductance(lcl["L2_grid_H"]), f"ripple gain at fsw = {v['ripple_gain_vs_L1_only']:.3f} (target ≤ {lcl['att_ratio_target']})"),
        ("Cf", format_capacitance(lcl["C_filter_F"]), "≤ 5 % base capacitance"),
        ("Rd", f"{lcl['R_damping_Ohm']:.3f} Ω", "passive damping"),
        ("Lf (LLCL trap)", format_inductance(lcl["Lf_trap_H"]) if lcl["Lf_trap_H"] else "n/a", "resonant at fsw"),
        ("Resonance frequency", f"{lcl['f_resonance_Hz']:.1f} Hz", f"status: {lcl['status']}"),
    ])
    for w in lcl.get("warnings", []):
        doc.add_paragraph(f"Warning: {w}")

    doc.add_heading("Chapter 3: Running the Simulation", level=1)
    notes = custom_notes.strip() if custom_notes else "none"
    doc.add_paragraph(
        f"Custom specifications: {notes}\n\n"
        "1. Extract the package and set the MATLAB current folder to it (Simscape Electrical required).\n"
        "2. Run run_simulation_workbench.m. It loads init_params.m, builds base_paper_circuit.slx and "
        "proposed_circuit.slx programmatically, simulates both, and prints grid-current THD.\n"
        "3. The models are averaged-inverter equivalents (fundamental + switching-ripple sources), "
        "fixed-step discrete solver, Ts = 1/(100·fsw).\n"
        "4. Compare the printed THD against the limit that applies to your point of common coupling.")

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out
